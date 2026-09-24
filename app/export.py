import html
import io
import os
import tempfile

from .translate import arabic_ratio, is_quranic

QURAN_FONT = "/usr/share/fonts/truetype/amiri-quran/AmiriQuran-Regular.ttf"


def _src_attrs(src):
    rtl = arabic_ratio(src) > 0.3
    cls = "src quran" if rtl and is_quranic(src) else "src"
    return cls, ("rtl" if rtl else "ltr"), ("ar" if rtl else "")

CSS = """
body { font-family: serif; line-height: 1.6; }
h2 { margin: 1.5em 0 .8em; }
p { margin: 0 0 .9em; text-align: justify; }
p.src { color: #666; font-size: .92em; margin-bottom: .3em; }
p.src[dir=rtl] { text-align: right; font-size: 1.05em; }
@font-face { font-family: "Amiri Quran"; src: url(fonts/AmiriQuran.ttf); }
p.src.quran { font-family: "Amiri Quran", serif; font-size: 1.25em; line-height: 2.1; }
"""


def _text(seg):
    return seg["out"] if seg["out"] is not None else seg["src"]


def _chapters(segments):
    chapters, cur_title, cur = [], None, []
    for s in segments:
        if s["kind"] == "h":
            if cur:
                chapters.append((cur_title, cur))
                cur = []
            cur_title = s
        else:
            cur.append(s)
    if cur or cur_title:
        chapters.append((cur_title, cur))
    return chapters


def to_txt(title, segments, bilingual):
    lines = [title, ""]
    for s in segments:
        if bilingual and s["out"] and s["out"] != s["src"]:
            lines.append(s["src"])
        lines.append(_text(s))
        lines.append("")
    return "\n".join(lines).encode("utf-8")


def to_epub(title, segments, bilingual, job_id):
    from ebooklib import epub

    book = epub.EpubBook()
    book.set_identifier(f"dedplay-ceviri-{job_id}")
    book.set_title(title)
    book.set_language("tr")
    style = epub.EpubItem(uid="style", file_name="style.css", media_type="text/css", content=CSS)
    book.add_item(style)
    if bilingual and os.path.exists(QURAN_FONT):
        book.add_item(epub.EpubItem(uid="amiri", file_name="fonts/AmiriQuran.ttf",
                                    media_type="font/ttf", content=open(QURAN_FONT, "rb").read()))

    items = []
    for n, (head, segs) in enumerate(_chapters(segments), 1):
        name = _text(head) if head else (title if n == 1 else f"Bölüm {n}")
        body = [f"<h2>{html.escape(name)}</h2>"]
        for s in segs:
            if bilingual and s["out"] and s["out"] != s["src"]:
                cls, d, lang = _src_attrs(s["src"])
                body.append(f'<p class="{cls}" dir="{d}">{html.escape(s["src"])}</p>')
            body.append(f"<p>{html.escape(_text(s))}</p>")
        ch = epub.EpubHtml(title=name[:80], file_name=f"bolum_{n:04d}.xhtml", lang="tr")
        ch.content = "<html><body>" + "\n".join(body) + "</body></html>"
        ch.add_item(style)
        book.add_item(ch)
        items.append(ch)

    book.toc = items
    book.add_item(epub.EpubNcx())
    book.add_item(epub.EpubNav())
    book.spine = ["nav"] + items
    with tempfile.NamedTemporaryFile(suffix=".epub", delete=False) as f:
        tmp = f.name
    try:
        epub.write_epub(tmp, book)
        return open(tmp, "rb").read()
    finally:
        os.remove(tmp)


def _rtl(paragraph):
    from docx.oxml import OxmlElement

    paragraph._p.get_or_add_pPr().append(OxmlElement("w:bidi"))
    for run in paragraph.runs:
        run._r.get_or_add_rPr().append(OxmlElement("w:rtl"))


def _quran_font(run):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    run.font.name = "Amiri Quran"
    rpr = run._r.get_or_add_rPr()
    fonts = rpr.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.insert(0, fonts)
    fonts.set(qn("w:cs"), "Amiri Quran")


def to_docx(title, segments, bilingual):
    import docx
    from docx.shared import Pt, RGBColor

    d = docx.Document()
    d.core_properties.title = title
    d.add_heading(title, level=0)
    for s in segments:
        if s["kind"] == "h":
            d.add_heading(_text(s), level=1)
            continue
        if bilingual and s["out"] and s["out"] != s["src"]:
            p = d.add_paragraph()
            r = p.add_run(s["src"])
            r.font.color.rgb = RGBColor(0x66, 0x66, 0x66)
            r.font.size = Pt(10.5)
            if arabic_ratio(s["src"]) > 0.3:
                if is_quranic(s["src"]):
                    _quran_font(r)
                    r.font.size = Pt(14)
                _rtl(p)
        d.add_paragraph(_text(s))
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


HTML_CSS = """
body { font-family: "Noto Serif", Georgia, serif; line-height: 1.7; color: #1c1c1c;
       max-width: 40em; margin: 0 auto; padding: 2em 1.2em 4em; background: #fdfdfb; }
.title { text-align: center; margin: 3em 0 4em; }
.title h1 { font-size: 2em; line-height: 1.25; margin: 0 0 .5em; }
.title p { color: #777; margin: 0; }
nav { margin-bottom: 4em; } nav ol { padding-left: 1.2em; } nav a { color: #2c4a7c; text-decoration: none; }
h2 { font-size: 1.4em; margin: 2.5em 0 1em; }
p { margin: 0 0 .9em; text-align: justify; hyphens: auto; }
p.src { color: #777; font-size: .92em; margin-bottom: .3em; }
p.src[dir=rtl] { font-family: "Noto Naskh Arabic", serif; font-size: 1.1em; text-align: right; }
p.src.quran { font-family: "Amiri Quran", "Noto Naskh Arabic", serif; font-size: 1.35em; line-height: 2.2; }
"""

PDF_CSS = """
@page { size: A5; margin: 18mm 15mm 20mm;
        @bottom-center { content: counter(page); font-family: "Noto Serif", serif; font-size: 8.5pt; color: #888; } }
@page :first { @bottom-center { content: none; } }
body { font-family: "Noto Serif", "Noto Naskh Arabic", serif; font-size: 10pt; line-height: 1.55; color: #111; }
.title { page-break-after: always; text-align: center; padding-top: 35%; }
.title h1 { font-size: 20pt; line-height: 1.3; margin: 0 0 12pt; }
.title p { color: #777; font-size: 10pt; margin: 0; }
nav { display: none; }
h2 { font-size: 14pt; margin: 18pt 0 12pt; bookmark-level: 1; }
h2.new-page { page-break-before: always; margin-top: 30pt; }
p { margin: 0 0 7pt; text-align: justify; hyphens: auto; orphans: 2; widows: 2; }
p.src { color: #666; font-size: 9pt; margin-bottom: 3pt; }
p.src[dir=rtl] { font-family: "Noto Naskh Arabic", serif; font-size: 10.5pt; text-align: right; }
p.src.quran { font-family: "Amiri Quran", "Noto Naskh Arabic", serif; font-size: 12.5pt; line-height: 2.1; }
"""


FONT_LINK = ('<link rel="stylesheet" href="https://fonts.googleapis.com/css2?'
             'family=Amiri+Quran&family=Noto+Naskh+Arabic&family=Noto+Serif:wght@400;700&display=swap">')


def to_html(title, segments, bilingual, css=HTML_CSS, head=FONT_LINK):
    chapters = _chapters(segments)
    body, toc = [], []
    for n, (head, segs) in enumerate(chapters, 1):
        if head:
            name = html.escape(_text(head))
            cls = "new-page" if len(segs) >= 5 else ""
            body.append(f'<h2 id="b{n}" class="{cls}">{name}</h2>')
            toc.append(f'<li><a href="#b{n}">{name}</a></li>')
        for s in segs:
            if bilingual and s["out"] and s["out"] != s["src"]:
                cls, d, lang = _src_attrs(s["src"])
                body.append(f'<p class="{cls}" dir="{d}" lang="{lang}">{html.escape(s["src"])}</p>')
            body.append(f"<p>{html.escape(_text(s))}</p>")
    nav = f"<nav><h2>İçindekiler</h2><ol>{''.join(toc)}</ol></nav>" if len(toc) > 1 else ""
    t = html.escape(title)
    doc = (f'<!doctype html><html lang="tr"><head><meta charset="utf-8">'
           f'<meta name="viewport" content="width=device-width, initial-scale=1">'
           f"<title>{t}</title>{head}<style>{css}</style></head><body>"
           f'<div class="title"><h1>{t}</h1><p>Türkçe çeviri</p></div>{nav}'
           + "\n".join(body) + "</body></html>")
    return doc.encode("utf-8")


def to_pdf(title, segments, bilingual):
    from weasyprint import HTML

    return HTML(string=to_html(title, segments, bilingual, css=PDF_CSS, head="").decode("utf-8")).write_pdf()


def build(job, segments, fmt, bilingual):
    title = job["title"] or os.path.splitext(job["filename"])[0]
    partial = job["status"] != "done"
    base = os.path.splitext(job["filename"])[0] + ("_tr_iki-dilli" if bilingual else "_tr")
    if partial:
        base += "_kismi"
        segments = [s for s in segments if s["out"] is not None]
    if fmt == "epub":
        return to_epub(title, segments, bilingual, job["id"]), "application/epub+zip", base + ".epub"
    if fmt == "docx":
        mt = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        return to_docx(title, segments, bilingual), mt, base + ".docx"
    if fmt == "pdf":
        return to_pdf(title, segments, bilingual), "application/pdf", base + ".pdf"
    if fmt == "html":
        return to_html(title, segments, bilingual), "text/html; charset=utf-8", base + ".html"
    return to_txt(title, segments, bilingual), "text/plain; charset=utf-8", base + ".txt"
